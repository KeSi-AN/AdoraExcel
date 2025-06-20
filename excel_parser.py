import os
import hashlib
import re
import openpyxl
from datetime import datetime
from typing import Dict, List, Any, Optional
from pathlib import Path
from serializers import serialize_data

def calculate_file_hash(file_content: bytes) -> str:
    """Calculate SHA-256 hash of file content."""
    return hashlib.sha256(file_content).hexdigest()

def extract_tables_from_sheet(sheet) -> Dict[str, List[Dict]]:
    """Extract tables from a single worksheet."""
    tables = {}
    
    # Process defined tables first
    for table in sheet.tables.values():
        ref = table.ref
        data = sheet[ref]
        rows = [[cell.value for cell in row] for row in data]
        if len(rows) < 2:  # Skip empty tables
            continue
            
        headers = [str(h) if h is not None else "" for h in rows[0]]
        table_data = [dict(zip(headers, (serialize_data(cell) for cell in row))) 
                     for row in rows[1:]]
        tables[table.name] = table_data

    # Process implicit tables (data between empty rows)
    all_values = list(sheet.values)
    table_starts = [0] + [i + 1 for i, row in enumerate(all_values) 
                         if all(cell is None for cell in row)]
    
    for i, start in enumerate(table_starts):
        end = table_starts[i + 1] - 1 if i + 1 < len(table_starts) else len(all_values)
        chunk = all_values[start:end]
        if len(chunk) < 2:  # Skip chunks that don't have at least a header and one row
            continue
            
        headers = [str(h) if h is not None else f"column_{i+1}" for i, h in enumerate(chunk[0])]
        table_data = [dict(zip(headers, (serialize_data(cell) for cell in row))) 
                     for row in chunk[1:] if any(cell is not None for cell in row)]
        
        if table_data:  # Only add non-empty tables
            table_name = f"Table_{len(tables) + 1}"
            tables[table_name] = table_data
    
    return tables

def extract_all_tables(file_path: str) -> Dict[str, Dict[str, List[Dict]]]:
    """Extract all tables from all sheets in an Excel file."""
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        all_tables = {}
        
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            tables = extract_tables_from_sheet(sheet)
            if tables:  # Only add sheets that have tables
                all_tables[sheet_name] = tables
                
        return all_tables
        
    except Exception as e:
        raise Exception(f"Error processing Excel file: {str(e)}")

def save_uploaded_file(uploaded_file, upload_folder: str) -> tuple:
    """Save uploaded file to disk and return its path and hash."""
    os.makedirs(upload_folder, exist_ok=True)
    
    # Create a safe filename
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r'[^\w.-]', '_', uploaded_file.name)
    unique_name = f"{timestamp}_{safe_name}"
    file_path = os.path.join(upload_folder, unique_name)
    
    # Save the file
    with open(file_path, "wb") as f:
        f.write(uploaded_file.getvalue())
    
    # Calculate file hash
    file_hash = calculate_file_hash(uploaded_file.getvalue())
    
    return file_path, file_hash
